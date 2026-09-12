import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


FILE_NAME = "student_results.xlsx"

if not os.path.exists(FILE_NAME):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"

    headers = [
        "Name", "Roll No.", "Class",
        "Subject 1", "Subject 2", "Subject 3",
        "Subject 4", "Subject 5",
        "Total Marks", "Percentage", "Result"
    ]

    sheet.append(headers)
    workbook.save(FILE_NAME)


def add_student():
    name = name_entry.get()
    roll = roll_entry.get()
    student_class = class_entry.get()

    try:
        marks = [
            int(sub1_entry.get()),
            int(sub2_entry.get()),
            int(sub3_entry.get()),
            int(sub4_entry.get()),
            int(sub5_entry.get())
        ]
    except ValueError:
        messagebox.showerror("Error", "Please enter valid marks.")
        return

    
    if name == "" or roll == "" or student_class == "":
        messagebox.showerror("Error", "Please fill all student details.")
        return

    
    if any(mark < 0 or mark > 100 for mark in marks):
        messagebox.showerror("Error", "Marks must be between 0 and 100.")
        return

    
    total = sum(marks)
    percentage = total / 5

   
    if all(mark >= 35 for mark in marks):
        result = "Pass"
    else:
        result = "Fail"

    
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == roll:
            messagebox.showerror("Error", "Roll No. already exists.")
            workbook.close()
            return

   
    sheet.append([
        name,
        roll,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        f"{percentage:.1f}%",
        result
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    messagebox.showinfo("Success", "Student record saved successfully.")

    
    clear_entries()




def clear_entries():
    name_entry.delete(0, tk.END)
    roll_entry.delete(0, tk.END)
    class_entry.delete(0, tk.END)

    sub1_entry.delete(0, tk.END)
    sub2_entry.delete(0, tk.END)
    sub3_entry.delete(0, tk.END)
    sub4_entry.delete(0, tk.END)
    sub5_entry.delete(0, tk.END)



def get_result():
    roll = result_roll_entry.get()

    if roll == "":
        messagebox.showerror("Error", "Please enter Roll No.")
        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[1]) == roll:

            found = True

            
            for item in result_tree.get_children():
                result_tree.delete(item)

            
            result_tree.insert("", tk.END, values=(
                row[0],    #Name   
                row[1],     #Roll no.   
                row[2],       
                row[8],       
                row[9],       
                row[10]       
            ))

            break

    workbook.close()

    if not found:
        messagebox.showerror("Not Found", "Student record not found.")

def show_all_results():


    for item in all_tree.get_children():
        all_tree.delete(item)

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):

        all_tree.insert("", tk.END, values=(
            row[0],      # Name
            row[1],      # Roll No.
            row[2],      # Class
            row[8],      # Total
            row[9],      # Percentage
            row[10]      # Result
        ))

    workbook.close()


root = tk.Tk()
root.title("Student Result Management System")
root.geometry("850x650")




title = tk.Label(
    root,
    text="STUDENT RESULT MANAGEMENT SYSTEM",
    font=("Arial", 20, "bold")
)

title.pack(pady=15)


add_frame = tk.LabelFrame(
    root,
    text="Add Student",
    font=("Arial", 12, "bold"),
    padx=15,
    pady=10
)

add_frame.pack(fill="x", padx=20, pady=10)


tk.Label(add_frame, text="Name").grid(row=0, column=0, padx=10, pady=5)
name_entry = tk.Entry(add_frame)
name_entry.grid(row=0, column=1, padx=10, pady=5)

tk.Label(add_frame, text="Roll No.").grid(row=0, column=2, padx=10, pady=5)
roll_entry = tk.Entry(add_frame)
roll_entry.grid(row=0, column=3, padx=10, pady=5)

tk.Label(add_frame, text="Class").grid(row=0, column=4, padx=10, pady=5)
class_entry = tk.Entry(add_frame)
class_entry.grid(row=0, column=5, padx=10, pady=5)


tk.Label(add_frame, text="Subject 1").grid(row=1, column=0, padx=10, pady=5)
sub1_entry = tk.Entry(add_frame, width=10)
sub1_entry.grid(row=1, column=1, padx=10, pady=5)

tk.Label(add_frame, text="Subject 2").grid(row=1, column=2, padx=10, pady=5)
sub2_entry = tk.Entry(add_frame, width=10)
sub2_entry.grid(row=1, column=3, padx=10, pady=5)

tk.Label(add_frame, text="Subject 3").grid(row=1, column=4, padx=10, pady=5)
sub3_entry = tk.Entry(add_frame, width=10)
sub3_entry.grid(row=1, column=5, padx=10, pady=5)


tk.Label(add_frame, text="Subject 4").grid(row=2, column=0, padx=10, pady=5)
sub4_entry = tk.Entry(add_frame, width=10)
sub4_entry.grid(row=2, column=1, padx=10, pady=5)

tk.Label(add_frame, text="Subject 5").grid(row=2, column=2, padx=10, pady=5)
sub5_entry = tk.Entry(add_frame, width=10)
sub5_entry.grid(row=2, column=3, padx=10, pady=5)




save_button = tk.Button(
    add_frame,
    text="💾 Save",
    command=add_student,
    width=15
)

save_button.grid(row=2, column=5, padx=10, pady=10)




result_frame = tk.LabelFrame(
    root,
    text="Get Result",
    font=("Arial", 12, "bold"),
    padx=10,
    pady=10
)

result_frame.pack(fill="x", padx=20, pady=10)


tk.Label(result_frame, text="Enter Roll No.:").pack(side="left", padx=10)

result_roll_entry = tk.Entry(result_frame)
result_roll_entry.pack(side="left", padx=10)

get_button = tk.Button(
    result_frame,
    text="🔍 Get Result",
    command=get_result
)

get_button.pack(side="left", padx=10)


result_columns = (
    "Name",
    "Roll No.",
    "Class",
    "Total Marks",
    "Percentage",
    "Result"
)

result_tree = ttk.Treeview(
    root,
    columns=result_columns,
    show="headings",
    height=3
)

for column in result_columns:
    result_tree.heading(column, text=column)
    result_tree.column(column, width=120)

result_tree.pack(fill="x", padx=20, pady=5)


all_frame = tk.LabelFrame(
    root,
    text="All Student Results",
    font=("Arial", 12, "bold"),
    padx=10,
    pady=10
)

all_frame.pack(fill="both", expand=True, padx=20, pady=10)


show_button = tk.Button(
    all_frame,
    text="📋 Show All Results",
    command=show_all_results
)

show_button.pack(pady=5)


all_tree = ttk.Treeview(
    all_frame,
    columns=result_columns,
    show="headings"
)

for column in result_columns:
    all_tree.heading(column, text=column)
    all_tree.column(column, width=120)

all_tree.pack(fill="both", expand=True)



exit_button = tk.Button(
    root,
    text="Exit",
    command=root.destroy,
    width=15
)

exit_button.pack(pady=10)


root.mainloop()
